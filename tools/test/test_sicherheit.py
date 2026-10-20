"""Regressionstests für die Grenze „fremder Anzeigentext → eigene Dateien“."""
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from tracker import als_text  # noqa: E402

try:
    from suche.quellen import portal  # braucht scrapling (tools/.venv)
except ImportError:  # pragma: no cover
    portal = None


class FormelTest(unittest.TestCase):
    def test_titel_mit_gleichheitszeichen_bleibt_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            datei = Path(tmp) / "t.xlsx"
            wb = openpyxl.Workbook()
            als_text(wb.active["A1"], '=HYPERLINK("https://example.invalid","x")')
            als_text(wb.active["A2"], "AI Consultant")
            wb.save(datei)
            ws = openpyxl.load_workbook(datei).active
            self.assertEqual(ws["A1"].data_type, "s")
            self.assertEqual(ws["A1"].value, '=HYPERLINK("https://example.invalid","x")')
            self.assertEqual(ws["A2"].value, "AI Consultant")


@unittest.skipIf(portal is None, "scrapling fehlt")
class PortalTest(unittest.TestCase):
    def test_nur_https_auf_bekannten_portalen(self):
        self.assertEqual(portal("https://www.stepstone.de/stellenangebote--x-1-inline.html"), "stepstone.de")
        self.assertEqual(portal("https://de.linkedin.com/jobs/view/x-4473165768"), "linkedin.com")

    def test_teilstring_reicht_nicht(self):
        self.assertIsNone(portal("https://fremd.example/?stepstone.de"))
        self.assertIsNone(portal("https://stepstone.de.fremd.example/"))
        self.assertIsNone(portal("http://www.stepstone.de/x"))
        self.assertIsNone(portal("javascript:alert(1)"))


if __name__ == "__main__":
    unittest.main()
