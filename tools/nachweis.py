#!/usr/bin/env python3
"""
Nachweis der Eigenbemühungen für die Agentur für Arbeit, aus dem Tracker.
Bewusst ohne Versanddaten und interne Notizen.

    tools/.venv/bin/python tools/nachweis.py
"""
import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from tracker import als_text  # noqa: E402

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

os.chdir(Path(__file__).resolve().parent.parent)
q = openpyxl.load_workbook('Bewerbungs-Tracker.xlsx')['Bewerbungen']
# Alles, was das Haus verlassen hat, zählt als versendet — auch mit Antwort.
VERSENDET = ('Gesendet', 'Eingangsbestätigung', 'Gespräch vereinbart', 'Im Gespräch', 'Angebot', 'Absage')


def hole(*status):
    d = [tuple(str(q[f'{sp}{r}'].value or '') for sp in 'CDEFKG')
         for r in range(3, 303) if q[f'O{r}'].value in status]
    return sorted(d, key=lambda x: x[0].lower())


versendet, vorbereitet = hole(*VERSENDET), hole('Vorbereitet')

wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Eigenbemühungen'
RAHMEN = Border(*[Side(style='thin', color='BBBBBB')]*4)
KOPF = ['Nr.', 'Unternehmen', 'Position', 'Ort / Modell', 'Gefunden über', 'Bewerbungsweg', 'Stellenanzeige']

ws['A1'] = 'Nachweis der Eigenbemühungen'; ws['A1'].font = Font(bold=True, size=16)
_p = person.person()
ws['A2'] = ' · '.join([person.name(), *str(_p.get('anschrift', '')).split(chr(10))]).strip(' ·')
ws['A3'] = ' · '.join(z for z in (_p['kontakt'][0] if isinstance(_p['kontakt'][0], list) else [_p['kontakt'][0]]))
for z in ('A2','A3'): ws[z].font = Font(size=10, color='555555')

def block(zeile, titel, daten, farbe, hinweis=None):
    ws.cell(row=zeile, column=1, value=titel).font = Font(bold=True, size=12)
    if hinweis:
        c = ws.cell(row=zeile+1, column=1, value=hinweis)
        c.font = Font(size=9, italic=True, color='777777')
        zeile += 1
    kz = zeile + 1
    for i, k in enumerate(KOPF, 1):
        c = ws.cell(row=kz, column=i, value=k)
        c.font = Font(bold=True, color='FFFFFF', size=10)
        c.fill = PatternFill('solid', fgColor=farbe); c.border = RAHMEN
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for n, reihe in enumerate(daten, 1):
        r = kz + n
        for i, wert in enumerate([n, *reihe], 1):
            c = ws.cell(row=r, column=i)
            als_text(c, wert)
            c.border = RAHMEN
            c.alignment = Alignment(vertical='top', wrap_text=(i in (2,3,4)),
                                    horizontal='center' if i == 1 else 'left')
            c.font = Font(size=9 if i == 7 else 10, color='0563C1' if i == 7 else '000000')
        ws.row_dimensions[r].height = 30
    return kz + len(daten) + 1

z = block(5, f'A. Versendete Bewerbungen ({len(versendet)})', versendet, '2F3E4E')
z = block(z + 2, f'B. Vorbereitet, noch nicht versendet ({len(vorbereitet)})', vorbereitet, '7A6A4F',
          'Unterlagen liegen vollständig vor. Diese Stellen zählen nicht als versendete Bewerbungen.')

ende = z + 3
ws.cell(row=ende, column=1, value='Ort, Datum').font = Font(size=10, color='555555')
ws.cell(row=ende, column=4, value='Unterschrift').font = Font(size=10, color='555555')
for sp in (1, 4):
    for off in (0, 1):
        ws.cell(row=ende-1, column=sp+off).border = Border(bottom=Side(style='thin', color='555555'))

for sp, b in zip('ABCDEFG', [5, 34, 42, 24, 14, 18, 58]): ws.column_dimensions[sp].width = b
ws.print_area = f'A1:G{ende}'
ws.page_setup.orientation = 'landscape'
ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.save('Nachweis-Eigenbemuehungen.xlsx')

print(f'A. Versendet ({len(versendet)}) · B. Vorbereitet ({len(vorbereitet)}) · gesamt {len(versendet)+len(vorbereitet)}\n')
print('B — vorbereitet, noch nicht versendet:')
for n, d in enumerate(vorbereitet, 1): print(f'  {n}. {d[0][:42]:<44} {d[1][:40]}')
