#!/usr/bin/env python3
"""
Bewerbungs-Tracker (Bewerbungs-Tracker.xlsx), Blatt „Bewerbungen“.

    tools/.venv/bin/python tools/tracker.py liste
    tools/.venv/bin/python tools/tracker.py neu stellen/<ordner>
    tools/.venv/bin/python tools/tracker.py status <firma> <status> [TT.MM.JJJJ]

Spalten A (Nr), Q (Reaktion) und R (Nachfassen am) sind Formeln und werden nie
überschrieben. Vor jedem Speichern liegt eine Sicherung als .bak daneben.
"""
from __future__ import annotations

import datetime as dt
import re
import shutil
import sys
from pathlib import Path

WURZEL = Path(__file__).resolve().parent.parent
DATEI = WURZEL / "Bewerbungs-Tracker.xlsx"
ERSTE_ZEILE = 3
LETZTE_ZEILE = 302
STATUS = ("Vorbereitet", "Gesendet", "Eingangsbestätigung", "Gespräch vereinbart",
          "Im Gespräch", "Angebot", "Absage", "Zurückgezogen")
SPALTEN = {
    "datum": "B", "firma": "C", "rolle": "D", "ort": "E", "quelle": "F", "url": "G",
    "ansprechpartner": "H", "email": "I", "telefon": "J", "kanal": "K", "passung": "L",
    "gehalt": "M", "unterlagen": "N", "status": "O", "antwort": "P", "nachgefasst": "S",
    "naechster_schritt": "T", "notizen": "U",
}


RECHTSFORM = re.compile(r"\b(gmbh|ag|se|kg|kgaa|co|mbh|ug|ev|ltd|inc|llc|plc|sa|holding|group|gruppe|"
                        r"deutschland|germany|international)\b")


def norm_firma(name: str) -> str:
    """Vergleichsschlüssel: klein, Umlaute aufgelöst, ohne Rechtsform, nur a–z/0–9."""
    t = (name or "").lower()
    for a, b in (("ä", "ae"), ("ö", "oe"), ("ü", "ue"), ("ß", "ss")):
        t = t.replace(a, b)
    t = RECHTSFORM.sub(" ", re.sub(r"[&.,/()+-]", " ", t))
    return re.sub(r"[^a-z0-9]", "", t)


def gleiche_firma(a: str, b: str) -> bool:
    """Gleich, wenn die Schlüssel gleich sind oder einer (ab 6 Zeichen) Präfix des anderen ist."""
    x, y = norm_firma(a), norm_firma(b)
    if len(x) < 3 or len(y) < 3:
        return False
    if x == y:
        return True
    kurz, lang = sorted((x, y), key=len)
    return len(kurz) >= 6 and lang.startswith(kurz)


def stelle_lesen(ordner: Path) -> dict:
    """Kopfdaten aus stelle.md und die Passung aus analyse.md eines Stellenordners."""
    daten: dict = {"unterlagen": f"stellen/{ordner.name}"}
    stelle = ordner / "stelle.md"
    if stelle.exists():
        zeilen = stelle.read_text(encoding="utf-8").splitlines()
        if zeilen and zeilen[0].startswith("# "):
            daten["rolle"] = zeilen[0][2:].strip()
        for z in zeilen[1:6]:
            m = re.match(r"\*\*(.+?)\*\*\s*·\s*([^·]*)", z)
            if m:
                daten["firma"] = m.group(1).strip()
                daten["ort"] = re.sub(r"\s+\d{5}\b", "", m.group(2)).strip()
                break
        for z in zeilen:
            m = re.match(r"\|\s*(Quelle|URL)\s*\|\s*(.+?)\s*\|", z)
            if m:
                daten["quelle" if m.group(1) == "Quelle" else "url"] = m.group(2)
    analyse = ordner / "analyse.md"
    if analyse.exists():
        m = re.search(r"Passung (\d+) von 100", analyse.read_text(encoding="utf-8"))
        if m:
            daten["passung"] = int(m.group(1))
    return daten


FORMEL_ANFANG = ("=", "+", "-", "@", "\t", "\r")


def als_text(zelle, wert) -> None:
    """Schreibt einen Wert; Text aus Anzeigen wird nie zur Formel (openpyxl macht aus „=…“ sonst eine)."""
    zelle.value = wert
    if isinstance(wert, str) and wert.startswith(FORMEL_ANFANG):
        zelle.data_type = "s"


def _datum(text: str | None):
    if not text:
        return None
    return dt.datetime.strptime(text, "%d.%m.%Y")


class Tracker:
    def __init__(self, datei: Path = DATEI):
        import openpyxl  # erst hier, damit die reinen Funktionen ohne openpyxl testbar bleiben

        self.datei = datei
        self.wb = openpyxl.load_workbook(datei)
        self.ws = self.wb["Bewerbungen"]

    def zeilen(self):
        for r in range(ERSTE_ZEILE, LETZTE_ZEILE + 1):
            firma = self.ws[f"C{r}"].value
            if firma:
                yield r, {k: self.ws[f"{s}{r}"].value for k, s in SPALTEN.items()}

    def finden(self, firma: str) -> int | None:
        treffer = [r for r, z in self.zeilen() if gleiche_firma(str(z["firma"]), firma)]
        if len(treffer) > 1:
            raise ValueError(f"„{firma}“ ist mehrdeutig (Zeilen {treffer})")
        return treffer[0] if treffer else None

    def freie_zeile(self) -> int:
        for r in range(ERSTE_ZEILE, LETZTE_ZEILE + 1):
            if not self.ws[f"C{r}"].value:
                return r
        raise RuntimeError("Tracker ist voll (Zeile 302 erreicht)")

    def eintragen(self, daten: dict, zeile: int | None = None) -> int:
        unbekannt = set(daten) - set(SPALTEN)
        if unbekannt:
            raise KeyError(f"Unbekannte Felder: {sorted(unbekannt)}")
        if daten.get("status") and daten["status"] not in STATUS:
            raise ValueError(f"Status „{daten['status']}“ nicht erlaubt: {STATUS}")
        r = zeile or self.freie_zeile()
        for k, v in daten.items():
            als_text(self.ws[f"{SPALTEN[k]}{r}"], v)
        return r

    def speichern(self) -> None:
        if self.datei.exists():
            shutil.copy2(self.datei, self.datei.with_suffix(".xlsx.bak"))
        self.wb.save(self.datei)


def befehl_liste(t: Tracker) -> None:
    for r, z in t.zeilen():
        datum = z["datum"].strftime("%d.%m.") if isinstance(z["datum"], dt.datetime) else "      "
        print(f"{r:>3}  {datum}  {str(z['status'] or ''):<20} {str(z['passung'] or ''):>3}  "
              f"{str(z['firma'])[:34]:<35} {str(z['rolle'] or '')[:50]}")


def befehl_neu(t: Tracker, ordner: str) -> None:
    pfad = (WURZEL / ordner).resolve()
    if not pfad.is_dir() or WURZEL / "stellen" not in pfad.parents:
        raise SystemExit(f"Kein Stellenordner unter stellen/: {ordner}")
    daten = stelle_lesen(pfad)
    if "firma" not in daten:
        raise SystemExit(f"{ordner}/stelle.md nennt keine Firma (Zeile „**Firma** · Ort“)")
    if t.finden(daten["firma"]):
        raise SystemExit(f"{daten['firma']} steht schon im Tracker — keine Dublette angelegt")
    r = t.eintragen({**daten, "status": "Vorbereitet"})
    t.speichern()
    print(f"Zeile {r}: {daten['firma']} — {daten.get('rolle', '')} (Vorbereitet)")


def befehl_status(t: Tracker, firma: str, status: str, datum: str | None) -> None:
    r = t.finden(firma)
    if r is None:
        raise SystemExit(f"„{firma}“ steht nicht im Tracker")
    daten: dict = {"status": status}
    tag = _datum(datum) or dt.datetime.combine(dt.date.today(), dt.time())
    if status == "Gesendet":
        daten["datum"] = tag
    elif status in ("Eingangsbestätigung", "Gespräch vereinbart", "Absage", "Angebot") and not t.ws[f"P{r}"].value:
        daten["antwort"] = tag
    t.eintragen(daten, zeile=r)
    t.speichern()
    print(f"Zeile {r}: {t.ws[f'C{r}'].value} → {status}")


def main(argv: list[str]) -> None:
    if not argv or argv[0] not in ("liste", "neu", "status"):
        raise SystemExit(__doc__)
    t = Tracker()
    if argv[0] == "liste":
        befehl_liste(t)
    elif argv[0] == "neu" and len(argv) == 2:
        befehl_neu(t, argv[1])
    elif argv[0] == "status" and len(argv) in (3, 4):
        befehl_status(t, argv[1], argv[2], argv[3] if len(argv) == 4 else None)
    else:
        raise SystemExit(__doc__)


if __name__ == "__main__":
    main(sys.argv[1:])
